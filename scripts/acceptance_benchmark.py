"""Bounded, reproducible acceptance measurements; all generated data stays in the target."""

import argparse
import json
import math
import os
import platform
import statistics
import threading
import time
from pathlib import Path
from types import SimpleNamespace

import psutil

from shiwen import __version__
from shiwen.library import Library
from shiwen.store import Store


def summary(values):
    values = sorted(values)
    return {
        "samples": len(values),
        "median_ms": statistics.median(values),
        "p95_ms": values[math.ceil(len(values) * 0.95) - 1],
        "max_ms": values[-1],
    }


def generate(folder, count):
    from docx import Document
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen.canvas import Canvas

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    folder.mkdir(parents=True)
    formats = {"pdf": 0, "md": 0, "docx": 0}
    for index in range(count):
        kind = "pdf" if index % 10 < 4 else "md" if index % 10 < 8 else "docx"
        formats[kind] += 1
        lines = [f"Record{index:06d} Topic{index % 100:03d} 离线部署 offline deployment"]
        lines += [
            f"Section {part}: 本地资料与检索记录。Budget review, project notes {index * 31 + part}."
            for part in range(80)
        ]
        path = folder / f"Report-{index:06d}.{kind}"
        if kind == "md":
            path.write_text("# 项目资料\n\n" + "\n".join(lines), encoding="utf-8")
        elif kind == "docx":
            document = Document()
            document.add_heading("项目资料 / Project notes", 0)
            for line in lines:
                document.add_paragraph(line)
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "负责人 / Owner"
            table.cell(1, 1).text = "Local review"
            document.save(path)
        else:
            canvas = Canvas(str(path), invariant=1)
            for start in range(0, len(lines), 40):
                canvas.setFont("STSong-Light", 10)
                for row, line in enumerate(lines[start : start + 40]):
                    canvas.drawString(35, 795 - row * 17, line)
                canvas.showPage()
            canvas.save()
        os.utime(path, (1700000000, 1700000000))
    return formats


def query_suite(library, count):
    timings = {"topic": [], "unique": [], "phrase_and": [], "absent": []}
    for index in range(100):
        group = index % 4
        number = index % min(count, 100)
        kind, query, expected = (
            ("topic", f"Topic{number:03d}", min(50, len(range(number, count, 100)))),
            ("unique", f"Record{index % count:06d}", 1),
            (
                "phrase_and",
                f'"offline deployment" Topic{number:03d}',
                min(50, len(range(number, count, 100))),
            ),
            ("absent", f"不存在的验收词{index:03d}", 0),
        )[group]
        before = time.perf_counter()
        result = library.search(query=query, mode="content")
        timings[kind].append((time.perf_counter() - before) * 1000)
        assert len(result["items"]) == expected, (kind, index, expected, len(result["items"]))
    return {key: summary(values) for key, values in timings.items()}


def metadata_test(directory, count):
    store = Store(directory)
    try:
        before = time.perf_counter()
        for start in range(0, count, 128):
            entries = []
            for index in range(start, min(start + 128, count)):
                path = directory / "virtual" / f"group-{index % 100}" / f"File-{index:07d}.txt"
                stat = SimpleNamespace(st_size=1024, st_mtime_ns=index)
                entries.append((path, stat, str(index), "txt", "metadata"))
            store.discover_batch(entries, "benchmark")
        build = time.perf_counter() - before
        timings = {}
        for query in ["", "File", f"{count - 1:07d}", "f", "notfound"]:
            values = []
            for _ in range(5):
                before = time.perf_counter()
                result = store.search(query, mode="name")
                values.append((time.perf_counter() - before) * 1000)
                assert len(result["items"]) == (
                    0
                    if query == "notfound"
                    else 1
                    if query == f"{count - 1:07d}"
                    else min(50, count)
                )
            timings[query or "empty"] = summary(values)
        return {
            "entries": count,
            "build_seconds": build,
            "queries": timings,
            "index_bytes": sum(p.stat().st_size for p in directory.glob("library.sqlite3*")),
            "scope": "Synthetic metadata only, no real filesystem permission checks or traversal",
        }
    finally:
        store.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("directory", type=Path)
    parser.add_argument("--documents", type=int, default=1000)
    parser.add_argument("--metadata", type=int, default=0)
    args = parser.parse_args()
    if args.documents < 100 or args.metadata < 0:
        parser.error("At least 100 documents; metadata must be nonnegative")
    args.directory.mkdir(parents=True, exist_ok=False)
    report = {
        "version": __version__,
        "platform": platform.platform(),
        "python": platform.python_version(),
        "logical_cpus": psutil.cpu_count(),
        "memory_bytes": psutil.virtual_memory().total,
        "limitations": [
            "Generated corpus, not private real-world documents",
            "No OS cache eviction: reopen measurements are not cold-cache results",
            "Benchmark-controller and parser RSS; desktop measured separately",
        ],
    }

    def save():
        (args.directory / "report.json").write_text(json.dumps(report, indent=2) + "\n")

    folder = args.directory / "documents"
    report["corpus"] = {
        "count": args.documents,
        "formats": generate(folder, args.documents),
        "file_bytes": sum(p.stat().st_size for p in folder.iterdir()),
        "docx": "Complete OOXML generated with python-docx; paragraphs and tables",
        "pdf": "CJK text PDFs, three pages each",
    }
    save()
    print("Generated corpus", report["corpus"], flush=True)
    process = psutil.Process()
    stopped = threading.Event()
    samples = []

    def monitor():
        while not stopped.wait(0.1):
            total = 0
            for child in [process, *process.children(recursive=True)]:
                try:
                    total += child.memory_info().rss
                except psutil.Error:
                    pass
            samples.append(total)

    watcher = threading.Thread(target=monitor, daemon=True)
    watcher.start()
    library = Library(
        args.directory / "index", automatic=True, disk_provider=lambda: ([folder], set())
    )
    try:
        before = time.perf_counter()
        library.scan()
        report["name_discovery_seconds"] = time.perf_counter() - before
        assert library.store.counts()["total"] == args.documents
        assert len(library.search(query="Report", mode="name")["items"]) == 50
        name_samples = []
        worker = threading.Thread(target=library.parse_pending)
        before = time.perf_counter()
        worker.start()
        while worker.is_alive():
            query_start = time.perf_counter()
            library.search(query="Report", mode="name")
            name_samples.append((time.perf_counter() - query_start) * 1000)
            worker.join(0.5)
        report["content_index_seconds"] = time.perf_counter() - before
        assert library.store.counts()["by_status"] == {"ready": args.documents}
        report["name_search_during_parsing"] = summary(name_samples)
        report["queries"] = query_suite(library, args.documents)
        report["index_bytes"] = library.status()["disk_bytes"]
        save()
        print("Document indexing and 100 queries passed", flush=True)
    finally:
        library.close()
    reopened = Library(
        args.directory / "index", automatic=True, disk_provider=lambda: ([folder], set())
    )
    try:
        before = time.perf_counter()
        assert reopened.search(query="Record000000")["items"]
        report["reopened_first_query_ms"] = (time.perf_counter() - before) * 1000
    finally:
        reopened.close()
    if args.metadata:
        report["metadata"] = metadata_test(args.directory / "metadata-index", args.metadata)
    stopped.set()
    watcher.join()
    report["controller_and_parser_peak_rss_bytes"] = max(samples, default=0)
    report["ok"] = True
    save()
    print(json.dumps(report, indent=2), flush=True)


if __name__ == "__main__":
    main()
